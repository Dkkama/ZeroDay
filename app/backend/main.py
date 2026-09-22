from __future__ import annotations

import hashlib
import hmac
import json
import mimetypes
import threading
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.exceptions import HTTPException as StarletteHTTPException
from pydantic import BaseModel

from config import (
    CATEGORY_META,
    COMPARE_FIELDS,
    CRON_SECRET,
    DATA_V2,
    DEMO_EMAIL,
    DEMO_PASSWORD,
    FLASH_SCORE,
    SESSION_SECRET,
    UPLOAD_DIR,
)
from exports import (
    audit_to_csv,
    audit_to_pdf,
    audit_to_txt,
    filter_emails,
    preset_spec,
    results_rows,
    submission_payload,
    to_csv,
    to_json,
    to_xlsx,
)
from ingest import _find_inbox_root, fetch_imap, ingest_zip
from llm import apply_decision, classify, classify_cursor_inbox, classify_vertex_batch, guard_vertex_batch
from seed import (
    build_record,
    load_attachment,
    load_flash_score,
    load_one_attached,
    next_live_id,
    seed_sample,
)
from firestore_store import open_store
from store import public_email

store = open_store()
app = FastAPI(title="ZeroDay SDOC", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173", "*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class LoginBody(BaseModel):
    email: str
    password: str


class SettingsBody(BaseModel):
    llm_provider: Optional[str] = None
    imap: Optional[dict] = None


class FieldEdit(BaseModel):
    name: str
    source: Optional[str] = None
    value: Optional[str] = None


class ProcessBody(BaseModel):
    email_ids: Optional[list[str]] = None
    all_unprocessed: bool = False


def _token() -> str:
    return hmac.new(SESSION_SECRET.encode(), DEMO_EMAIL.encode(), hashlib.sha256).hexdigest()


def require_auth(authorization: Optional[str]):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "sign in required")
    if not hmac.compare_digest(authorization.split(" ", 1)[1], _token()):
        raise HTTPException(401, "invalid session")


@app.post("/api/login")
def login(body: LoginBody):
    if body.email.strip().lower() != DEMO_EMAIL.lower() or body.password != DEMO_PASSWORD:
        raise HTTPException(401, "unknown clerk")
    return {"token": _token(), "email": DEMO_EMAIL, "name": "Clerk"}


@app.get("/api/me")
def me(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    return {"email": DEMO_EMAIL, "name": "Clerk"}


@app.get("/api/meta")
def meta(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    score = load_flash_score()
    final = None
    try:
        final = score.get("full", {}).get("final_score")
    except Exception:
        final = None
    return {
        "categories": CATEGORY_META,
        "fields": COMPARE_FIELDS,
        "demo_score": final,
        "score_path": str(FLASH_SCORE),
    }


@app.post("/api/seed")
def seed(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    n = seed_sample(store, replace=True)
    return {"loaded": n}


@app.get("/api/settings")
def get_settings(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    settings = store.settings()
    imap = dict(settings.get("imap") or {})
    if imap.get("password"):
        imap["password"] = "••••••"
    settings["imap"] = imap
    return settings


@app.put("/api/settings")
def put_settings(body: SettingsBody, authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    return store.update_settings(body.model_dump(exclude_none=True))


_seed_lock = threading.Lock()


def ensure_seed():
    if store.settings().get("auto_seed") is False:
        return
    if store.all_emails():
        return
    with _seed_lock:
        if store.all_emails():
            return
        seed_sample(store, replace=False)


@app.get("/api/stats")
def stats(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    ensure_seed()
    rows = store.all_emails()
    by_cat = {k: 0 for k in CATEGORY_META}
    by_status = {"OK": 0, "MISMATCH": 0, "NEEDS_REVIEW": 0}
    reasons = {}
    for rec in rows:
        by_cat[rec.get("category") or "GENERAL"] = by_cat.get(rec.get("category") or "GENERAL", 0) + 1
        by_status[rec.get("status") or "OK"] = by_status.get(rec.get("status") or "OK", 0) + 1
        if rec.get("review_reason"):
            reasons[rec["review_reason"]] = reasons.get(rec["review_reason"], 0) + 1
    queue = [r for r in rows if r.get("category") == "BL_COMPARISON"
             and r.get("status") in ("NEEDS_REVIEW", "MISMATCH")
             and not r.get("human_validated") and not r.get("fixed")]
    return {
        "total": len(rows),
        "by_category": by_cat,
        "by_status": by_status,
        "review_reasons": reasons,
        "check_documents": by_cat.get("BL_COMPARISON", 0),
        "needs_human": len(queue),
    }


@app.get("/api/emails")
def list_emails(
    authorization: Optional[str] = Header(None),
    category: Optional[str] = None,
    status: Optional[str] = None,
    q: Optional[str] = None,
    queue: bool = False,
):
    require_auth(authorization)
    ensure_seed()
    rows = store.all_emails()
    if queue:
        rows = [r for r in rows
                if r.get("category") == "BL_COMPARISON"
                and (r.get("status") in ("NEEDS_REVIEW", "MISMATCH") or r.get("fixed"))]
        rows.sort(key=lambda r: (
            1 if r.get("human_validated") or r.get("fixed") else 0,
            0 if r.get("status") == "NEEDS_REVIEW" else 1,
            r.get("email_id") or "",
        ))
    if category:
        rows = [r for r in rows if r.get("category") == category]
    if status:
        rows = [r for r in rows if r.get("status") == status]
    if q:
        needle = q.lower()
        rows = [r for r in rows
                if needle in (r.get("subject") or "").lower()
                or needle in (r.get("email_id") or "").lower()
                or needle in (r.get("from") or "").lower()]
    return [public_email(r) for r in rows]


@app.get("/api/emails/{email_id}")
def get_email(email_id: str, authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    rec = store.get_email(email_id)
    if not rec:
        raise HTTPException(404, "email not found")
    return public_email(rec, include_text=True)


def _file_bytes(email_id: str, att: dict) -> bytes | None:
    path = Path(att.get("abs") or "")
    if path.is_file():
        return path.read_bytes()
    name = att.get("filename")
    if not name:
        return None
    return store.get_file(email_id, name)


def _keep_files(rec: dict) -> None:
    for att in rec.get("attachments") or []:
        path = Path(att.get("abs") or "")
        name = att.get("filename")
        if name and path.is_file():
            store.put_file(rec["email_id"], name, path.read_bytes())


@app.get("/api/emails/{email_id}/attachments/{filename}")
def get_attachment(email_id: str, filename: str,
                   authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    rec = store.get_email(email_id)
    if not rec:
        raise HTTPException(404, "email not found")
    for att in rec.get("attachments") or []:
        if att.get("filename") == filename:
            data = _file_bytes(email_id, att)
            if data is not None:
                media = mimetypes.guess_type(filename)[0] or "application/octet-stream"
                return Response(
                    data,
                    media_type=media,
                    headers={"Content-Disposition": f'inline; filename="{filename}"'},
                )
            text = att.get("text") or ""
            return Response(text.encode("utf-8"), media_type="text/plain")
    raise HTTPException(404, "attachment not found")


@app.get("/api/emails/{email_id}/attachments/{filename}/sheet")
def get_sheet(email_id: str, filename: str,
              authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    rec = store.get_email(email_id)
    if not rec:
        raise HTTPException(404, "email not found")
    for att in rec.get("attachments") or []:
        if att.get("filename") != filename:
            continue
        data = _file_bytes(email_id, att)
        if data is None:
            raise HTTPException(404, "file missing")
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if c is None else str(c) for c in row])
            sheets.append({"name": ws.title, "rows": rows})
        wb.close()
        return {"filename": filename, "sheets": sheets}
    raise HTTPException(404, "attachment not found")


@app.post("/api/emails/{email_id}/validate")
def validate(email_id: str, authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    current = store.get_email(email_id)
    if not current:
        raise HTTPException(404, "email not found")
    issue = current.get("status") in ("MISMATCH", "NEEDS_REVIEW")
    rec = store.update_email(
        email_id,
        {"human_validated": True, "fixed": issue or bool(current.get("fixed"))},
        actor="human",
        change_type="fixed by human",
        details={
            "human_validated": True,
            "fixed": True,
            "status": current.get("status"),
            "review_reason": current.get("review_reason"),
            "defect_fields": current.get("defect_fields"),
        },
    )
    return public_email(rec)


@app.post("/api/emails/{email_id}/fields")
def edit_fields(email_id: str, body: list[FieldEdit],
                authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    rec = store.get_email(email_id)
    if not rec:
        raise HTTPException(404, "email not found")
    resolved = dict(rec.get("resolved_fields") or {})
    changed = []
    si = rec.get("si_fields") or {}
    bl = rec.get("bl_fields") or {}
    for item in body:
        name = item.name
        if name not in COMPARE_FIELDS:
            continue
        if item.source == "si":
            value = si.get(name) or ""
        elif item.source == "bl":
            value = bl.get(name) or ""
        else:
            value = item.value or ""
        resolved[name] = value
        changed.append({"field": name, "source": item.source, "value": value})
    rec = store.update_email(
        email_id,
        {"resolved_fields": resolved, "human_validated": False},
        actor="human",
        change_type="edited fields: " + ", ".join(c["field"] for c in changed),
        details={"fields": changed},
    )
    return public_email(rec)


@app.get("/api/audit")
def audit(authorization: Optional[str] = Header(None), limit: int = 500):
    require_auth(authorization)
    return store.audit(limit=limit)


@app.get("/api/jobs")
def jobs(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    return store.jobs()


@app.post("/api/ingest/sample")
def ingest_sample(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    n = seed_sample(store, replace=True)
    return {"loaded": n, "source": "sample"}


@app.post("/api/ingest/upload")
def ingest_upload(file: UploadFile = File(...),
                  authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    data = file.file.read()
    try:
        root, emails = ingest_zip(data, file.filename or "upload.zip")
    except (FileNotFoundError, json.JSONDecodeError, UnicodeDecodeError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, f"Could not read that zip: {exc}")
    provider = store.settings()["llm_provider"]
    job = store.add_job("upload", provider, len(emails))
    store.update_job(job["id"], root=str(root))
    threading.Thread(target=_run_ingest, args=(job["id"], root, emails, "upload", provider),
                     daemon=True).start()
    return {"job": job, "count": len(emails)}


class ImapBody(BaseModel):
    limit: int = 15


@app.post("/api/ingest/imap/test")
def imap_test(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    cfg = store.settings()["imap"]
    try:
        fetch_imap(cfg, limit=1)
        store.set_imap_status("connected")
        return {"ok": True}
    except Exception as exc:
        store.set_imap_status("error")
        raise HTTPException(400, str(exc))


def _ingest_imap_new(limit: int = 15, sync: bool = True) -> dict:
    cfg = store.settings()["imap"]
    root, emails = fetch_imap(cfg, limit=limit)
    fresh = [e for e in emails if not store.get_email(e["email_id"])]
    store.set_imap_status("connected", datetime.now(timezone.utc).isoformat())
    if not fresh:
        return {"job": None, "count": 0}
    provider = store.settings()["llm_provider"]
    job = store.add_job("imap", provider, len(fresh))
    if sync:
        _run_ingest(job["id"], root, fresh, "imap", provider)
        jobs = [j for j in store.jobs() if j["id"] == job["id"]]
        return {"job": jobs[0] if jobs else job, "count": len(fresh)}
    threading.Thread(target=_run_ingest, args=(job["id"], root, fresh, "imap", provider),
                     daemon=True).start()
    return {"job": job, "count": len(fresh)}


@app.post("/api/ingest/imap/fetch")
def imap_fetch(body: ImapBody, authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    try:
        return _ingest_imap_new(limit=body.limit, sync=True)
    except Exception as exc:
        store.set_imap_status("error")
        raise HTTPException(400, str(exc))


@app.post("/api/cron/imap")
def cron_imap(x_cron_secret: Optional[str] = Header(None)):
    if not CRON_SECRET or x_cron_secret != CRON_SECRET:
        raise HTTPException(401, "bad cron secret")
    try:
        return _ingest_imap_new(limit=15, sync=True)
    except Exception as exc:
        store.set_imap_status("error")
        raise HTTPException(400, str(exc))


@app.post("/api/ingest/one")
def ingest_one(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    rec = load_one_attached(store)
    return {
        "email_id": rec["email_id"],
        "subject": rec.get("subject"),
        "attachments": [a.get("filename") or Path(a.get("rel") or "").name
                        for a in rec.get("attachments") or []],
        "source": rec.get("source"),
    }


@app.post("/api/process/one")
def process_one(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    provider = store.settings()["llm_provider"]
    try:
        guard_vertex_batch(provider, 1)
    except RuntimeError as exc:
        raise HTTPException(400, str(exc))
    rec = load_one_attached(store, next_live_id(store))
    job = store.add_job("process", provider, 1)
    _run_process(job["id"], [rec["email_id"]], provider)
    rec = store.get_email(rec["email_id"]) or rec
    jobs = [j for j in store.jobs() if j["id"] == job["id"]]
    return {
        "job": jobs[0] if jobs else job,
        "email_id": rec["email_id"],
        "provider": provider,
        "category": rec.get("category"),
        "status": rec.get("status"),
    }


@app.post("/api/process")
def process(body: ProcessBody, authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    provider = store.settings()["llm_provider"]
    if body.email_ids:
        ids = body.email_ids
    elif body.all_unprocessed:
        ids = [r["email_id"] for r in store.all_emails() if r.get("source") != "seed"
               and r.get("category") == "GENERAL" and not r.get("subject")]
        ids = [r["email_id"] for r in store.all_emails() if r.get("source") in ("upload", "imap")]
    else:
        raise HTTPException(400, "email_ids required")
    try:
        guard_vertex_batch(provider, len(ids))
    except RuntimeError as exc:
        raise HTTPException(400, str(exc))
    job = store.add_job("process", provider, len(ids))
    threading.Thread(target=_run_process, args=(job["id"], ids, provider), daemon=True).start()
    return {"job": job}


@app.post("/api/jobs/{job_id}/retry")
def retry_job(job_id: int, authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    provider = store.settings()["llm_provider"]
    failed_ids = []
    # retry last failed process by re-running all upload/imap emails without defects? 
    # simpler: re-run the same kind is not stored. Retry means process all non-seed pending.
    ids = [r["email_id"] for r in store.all_emails()
           if r.get("source") in ("upload", "imap") and not r.get("processed_live")]
    if not ids:
        ids = [r["email_id"] for r in store.all_emails() if r.get("source") in ("upload", "imap")]
    try:
        guard_vertex_batch(provider, len(ids))
    except RuntimeError as exc:
        raise HTTPException(400, str(exc))
    job = store.add_job("retry", provider, len(ids))
    threading.Thread(target=_run_process, args=(job["id"], ids, provider), daemon=True).start()
    return {"job": job}


@app.get("/api/export/results")
def export_results(
    authorization: Optional[str] = Header(None),
    fmt: str = "json",
    preset: Optional[str] = None,
    statuses: Optional[str] = None,
    categories: Optional[str] = None,
    human_validated: Optional[bool] = None,
):
    require_auth(authorization)
    spec = preset_spec(preset) if preset else {}
    if statuses:
        spec["statuses"] = [s.strip() for s in statuses.split(",") if s.strip()]
    if categories:
        spec["categories"] = [s.strip() for s in categories.split(",") if s.strip()]
    if human_validated is not None:
        spec["human_validated"] = human_validated
    rows = results_rows(filter_emails(store.all_emails(), spec))
    if fmt == "csv":
        return Response(to_csv(rows), media_type="text/csv",
                        headers={"Content-Disposition": "attachment; filename=results.csv"})
    if fmt == "xlsx":
        return Response(to_xlsx(rows),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": "attachment; filename=results.xlsx"})
    return Response(to_json(rows), media_type="application/json")


@app.get("/api/export/submission")
def export_submission(authorization: Optional[str] = Header(None)):
    require_auth(authorization)
    payload = submission_payload(store.all_emails())
    return Response(json.dumps(payload, indent=2).encode(), media_type="application/json",
                    headers={"Content-Disposition": "attachment; filename=submission.json"})


@app.get("/api/export/audit")
def export_audit(
    authorization: Optional[str] = Header(None),
    fmt: str = "json",
    limit: int = 200,
):
    require_auth(authorization)
    rows = store.audit(limit=limit)
    if fmt == "csv":
        return Response(audit_to_csv(rows), media_type="text/csv",
                        headers={"Content-Disposition": "attachment; filename=audit.csv"})
    if fmt == "xlsx":
        return Response(to_xlsx([
            {"id": r.get("id"), "email_id": r.get("email_id"), "actor": r.get("actor"),
             "change_type": r.get("change_type"), "category": r.get("category"),
             "created_at": r.get("created_at")}
            for r in rows
        ]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=audit.xlsx"})
    if fmt == "pdf":
        return Response(audit_to_pdf(rows), media_type="application/pdf",
                        headers={"Content-Disposition": "attachment; filename=audit.pdf"})
    if fmt == "txt":
        return Response(audit_to_txt(rows), media_type="text/plain",
                        headers={"Content-Disposition": "attachment; filename=audit.txt"})
    return rows


def _run_ingest(job_id: int, root, emails, source, provider, resume: bool = False):
    current = next((j for j in store.jobs() if j["id"] == job_id), None)
    done = (current or {}).get("done") or 0 if resume else 0
    failed = (current or {}).get("failed") or 0 if resume else 0
    store.update_job(job_id, status="running", done=done, failed=failed,
                     error=None, root=str(root))
    if provider == "cursor":
        _run_ingest_cursor(job_id, root, emails, source, done, failed)
        return
    try:
        guard_vertex_batch(provider, len(emails))
    except RuntimeError as exc:
        store.update_job(job_id, status="failed", error=str(exc))
        # still store raw emails without live classify
        from seed import build_record
        for email in emails:
            rec = build_record(email, {}, root, source=source)
            _keep_files(rec)
            store.upsert_email(rec, actor="program")
        return
    from seed import build_record
    group = []
    built = [build_record(email, {}, root, source=source) for email in emails]
    for rec in built:
        _keep_files(rec)
    for rec in built:
        group.append(rec)
        if len(group) == 8:
            done, failed = _flush_vertex_batch(job_id, group, done, failed, store_failures=True)
            group = []
    if group:
        done, failed = _flush_vertex_batch(job_id, group, done, failed, store_failures=True)
    store.update_job(job_id, status="done" if not failed else "failed", done=done, failed=failed)


def _run_ingest_cursor(job_id: int, root, emails, source, done: int, failed: int):
    from seed import build_record

    lock = threading.Lock()

    def on_each(email, decision, error):
        nonlocal done, failed
        rec = build_record(email, {}, root, source=source)
        _keep_files(rec)
        if error or not decision:
            with lock:
                failed += 1
                store.upsert_email(rec, actor="program")
                store.update_job(job_id, done=done, failed=failed,
                                 error=str(error or "classify failed")[:300])
            return
        rec.update(apply_decision(email, rec["attachments"], decision))
        rec["processed_live"] = True
        with lock:
            done += 1
            store.upsert_email(rec, actor="program")
            store.update_job(job_id, done=done, failed=failed)

    try:
        classify_cursor_inbox(emails, root, on_each)
    except Exception as exc:
        store.update_job(job_id, status="failed", error=str(exc)[:300], done=done, failed=failed)
        return
    store.update_job(job_id, status="done" if not failed else "failed", done=done, failed=failed)


def _resume_running_jobs():
    for job in store.jobs():
        if job.get("status") != "running" or job.get("provider") != "cursor":
            continue
        root = Path(job["root"]) if job.get("root") else None
        if root is None or not (root / "inbox").is_dir():
            root = _latest_inbox_root()
        if root is None:
            store.update_job(job["id"], status="failed", error="upload folder missing")
            continue
        emails = []
        for path in sorted((root / "inbox").glob("*.json")):
            if path.name.startswith("._"):
                continue
            emails.append(json.loads(path.read_text(encoding="utf-8")))
        pending = []
        for email in emails:
            rec = store.get_email(email["email_id"])
            if rec and rec.get("processed_live"):
                continue
            pending.append(email)
        threading.Thread(
            target=_run_ingest,
            args=(job["id"], root, pending, job.get("kind") or "upload", "cursor", True),
            daemon=True,
        ).start()


def _latest_inbox_root():
    if not UPLOAD_DIR.exists():
        return None
    dirs = sorted((p for p in UPLOAD_DIR.iterdir() if p.is_dir()),
                  key=lambda p: p.stat().st_mtime, reverse=True)
    for dest in dirs:
        try:
            return _find_inbox_root(dest)
        except FileNotFoundError:
            continue
    return None


def _run_process(job_id: int, ids: list[str], provider: str):
    store.update_job(job_id, status="running")
    if provider == "vertex" and len(ids) > 1:
        _run_process_vertex_batches(job_id, ids)
        return
    done = failed = 0
    for eid in ids:
        rec = store.get_email(eid)
        if not rec:
            failed += 1
            continue
        try:
            email = {k: rec.get(k) for k in ("email_id", "from", "subject", "body", "attachments")}
            email["attachments"] = [
                a.get("rel") or f"attachments/{a.get('filename')}"
                for a in rec.get("attachments") or []
            ]
            decision = classify(provider, rec, rec.get("attachments") or [])
            rec.update(decision)
            rec["processed_live"] = True
            rec["classified_by"] = provider
            store.upsert_email(rec, actor="program",
                               change_type=f'reclassified with {provider}')
            done += 1
        except Exception as exc:
            failed += 1
            store.update_job(job_id, error=str(exc))
        store.update_job(job_id, done=done, failed=failed)
    store.update_job(job_id, status="done" if not failed else "failed", done=done, failed=failed)


def _classify_vertex_chunk(recs: list[dict]) -> tuple[dict, str | None]:
    last = None
    for attempt in range(6):
        try:
            return classify_vertex_batch(recs), None
        except Exception as exc:
            last = exc
            msg = str(exc)
            if "429" in msg or "RESOURCE_EXHAUSTED" in msg:
                time.sleep(min(60, 2 ** attempt))
                continue
            if attempt == 0:
                time.sleep(1)
                continue
            break
    return {}, str(last)[:300] if last else "vertex batch failed"


def _save_vertex_batch(group: list[dict], store_failures: bool = False) -> tuple[int, int, str | None]:
    parsed, err = _classify_vertex_chunk(group)
    missing = [rec for rec in group if rec["email_id"] not in parsed]
    if missing and parsed:
        again, err2 = _classify_vertex_chunk(missing)
        parsed.update(again)
        err = err2 or err
    saved = missed = 0
    for rec in group:
        decision = parsed.get(rec["email_id"])
        if not decision:
            missed += 1
            if store_failures:
                store.upsert_email(rec, actor="program")
            continue
        rec.update(decision)
        rec["processed_live"] = True
        rec["classified_by"] = "vertex"
        store.upsert_email(rec, actor="program", change_type="reclassified with vertex")
        saved += 1
    return saved, missed, err


def _flush_vertex_batch(job_id: int, group: list[dict], done: int, failed: int,
                        store_failures: bool = False) -> tuple[int, int]:
    saved, missed, err = _save_vertex_batch(group, store_failures=store_failures)
    done += saved
    failed += missed
    if err and missed:
        store.update_job(job_id, done=done, failed=failed, error=err)
    else:
        store.update_job(job_id, done=done, failed=failed, error=None)
    return done, failed


def _run_process_vertex_batches(job_id: int, ids: list[str]):
    """Three workers, eight emails per prompt, five prompts a minute. Same pool as Cursor."""
    import queue
    import sys
    from config import ROOT
    sys.path.insert(0, str(ROOT / "sdoc_eval"))
    from run import RateLimiter  # noqa: E402

    limiter = RateLimiter([(5, 60.0), (280, 3600.0)])
    work = queue.Queue()
    for eid in ids:
        work.put(eid)
    lock = threading.Lock()
    state = {"done": 0, "failed": 0}
    workers = 3 if len(ids) > 8 else 1

    def worker():
        while True:
            group = []
            missing = 0
            for _ in range(8):
                try:
                    eid = work.get_nowait()
                except queue.Empty:
                    break
                rec = store.get_email(eid)
                if rec:
                    group.append(rec)
                else:
                    missing += 1
            if missing:
                with lock:
                    state["failed"] += missing
                    store.update_job(job_id, done=state["done"], failed=state["failed"])
            if not group:
                return
            limiter.wait()
            saved, missed, err = _save_vertex_batch(group)
            with lock:
                state["done"] += saved
                state["failed"] += missed
                if err and missed:
                    store.update_job(job_id, done=state["done"], failed=state["failed"], error=err)
                else:
                    store.update_job(job_id, done=state["done"], failed=state["failed"], error=None)

    threads = [threading.Thread(target=worker, name=f"vertex-w{i}") for i in range(workers)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    store.update_job(
        job_id,
        status="done" if not state["failed"] else "failed",
        done=state["done"],
        failed=state["failed"],
    )


def _imap_poll_loop():
    while True:
        try:
            secs = int(store.settings().get("imap", {}).get("poll_seconds") or 0)
        except Exception:
            secs = 0
        if secs < 30:
            time.sleep(10)
            continue
        try:
            _ingest_imap_new(limit=15, sync=True)
        except Exception:
            store.set_imap_status("error")
        time.sleep(secs)


@app.on_event("startup")
def startup():
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    threading.Thread(target=_imap_poll_loop, daemon=True).start()
    _resume_running_jobs()


class SPAStaticFiles(StaticFiles):
    """Client routes such as /login have no file. Reload should get the app."""

    async def get_response(self, path: str, scope):
        try:
            return await super().get_response(path, scope)
        except StarletteHTTPException as exc:
            if exc.status_code != 404 or Path(path).suffix or path.startswith("api"):
                raise
            return await super().get_response("index.html", scope)


FRONTEND_DIST = Path(__file__).resolve().parents[1] / "frontend" / "dist"
if FRONTEND_DIST.exists():
    app.mount("/", SPAStaticFiles(directory=FRONTEND_DIST, html=True), name="ui")
